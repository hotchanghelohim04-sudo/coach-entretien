import io
import json
from html import escape

import ollama
import streamlit as st
from docx import Document     
from pypdf import PdfReader    

MODEL = "gemma3:4b"   
NB_QUESTIONS = 10      
MAX_CHARS = 6000      

st.set_page_config(page_title="Coach d'entretien", page_icon="🎤")

CSS = """
<style>
.stApp {
    background: linear-gradient(135deg, #dbeafe 0%, #ede9fe 50%, #fce7f3 100%);
    background-attachment: fixed;
}
header[data-testid="stHeader"] {background: transparent;}
.block-container {max-width: 820px; padding-top: 2rem;}

.stApp p, .stApp li, .stApp label, .stApp h1, .stApp h2, .stApp h3 {color: #1e293b;}

.titre {
    font-size: 2.8rem; font-weight: 800; text-align: center; line-height: 1.1;
    background: linear-gradient(90deg, #4f46e5, #db2777);
    -webkit-background-clip: text; -webkit-text-fill-color: transparent;
}
.sous-titre {text-align: center; color: #475569; font-size: 1.1rem; margin-bottom: 1.5rem;}

.carte {
    background: rgba(255, 255, 255, 0.8);
    border-radius: 20px; padding: 1.2rem 1.5rem; margin-bottom: 1rem;
    box-shadow: 0 8px 30px rgba(99, 102, 241, 0.15);
    color: #1e293b;
}
.carte h4 {margin: 0 0 .4rem 0; color: #4f46e5;}
.carte ul {margin: 0; padding-left: 1.2rem;}

.stButton > button, .stDownloadButton > button {
    border-radius: 12px; border: 1px solid #c7d2fe; padding: .55rem 1.3rem;
    font-weight: 600; background: white; transition: all .15s;
}
.stButton > button p, .stDownloadButton > button p {color: #4f46e5;}
.stButton > button[kind="primary"], .stDownloadButton > button[kind="primary"] {
    background: linear-gradient(90deg, #6366f1, #8b5cf6); border: none;
}
.stButton > button[kind="primary"] p, .stDownloadButton > button[kind="primary"] p {color: white;}
.stButton > button:hover, .stDownloadButton > button:hover {
    transform: translateY(-2px); box-shadow: 0 6px 18px rgba(99, 102, 241, 0.35);
}

.stTextArea textarea {
    border-radius: 14px; background: rgba(255, 255, 255, 0.9); color: #1e293b;
}
[data-testid="stFileUploaderDropzone"] {
    background: rgba(255, 255, 255, 0.7); border-radius: 14px;
    border: 2px dashed #a5b4fc;
}
[data-testid="stFileUploaderDropzone"] * {color: #334155;}

[data-testid="stChatMessage"] {
    background: rgba(255, 255, 255, 0.85); border-radius: 18px;
    padding: 1rem; margin-bottom: .6rem;
    box-shadow: 0 4px 16px rgba(99, 102, 241, 0.12);
}
[data-testid="stChatMessage"] p {color: #1e293b;}
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)

if "ecran" not in st.session_state:
    st.session_state.ecran = "offre"   
    st.session_state.messages = []     
    st.session_state.feedback = None  



def lire_fichier(fichier):
    """Lit un PDF ou un Word (.docx) et renvoie son texte."""
    if fichier is None:
        return ""
    if fichier.name.lower().endswith(".pdf"):
        pages = PdfReader(fichier).pages
        return "\n".join(page.extract_text() or "" for page in pages)
    document = Document(fichier)
    return "\n".join(p.text for p in document.paragraphs)


def demander(messages, en_json=False):
    """Envoie la conversation à Gemma et renvoie sa réponse (texte)."""
    options = {"format": "json"} if en_json else {}
    reponse = ollama.chat(model=MODEL, messages=messages, **options)
    return reponse["message"]["content"]


def retour_accueil():
    """Bouton Quitter / Recommencer : on revient à la page d'accueil."""
    st.session_state.ecran = "offre"
    st.session_state.messages = []
    st.session_state.feedback = None


def demarrer(offre, cv):
    """Prépare le recruteur et pose la première question."""
    consigne = (
        "Tu es un recruteur qui fait passer un entretien de stage en français. "
        "Tu vouvoies le candidat.\n"
        f"Voici l'offre de stage :\n{offre[:MAX_CHARS]}\n\n"
    )
    if cv.strip():
        consigne += f"Voici le CV du candidat :\n{cv[:MAX_CHARS]}\n\n"
    consigne += (
        "Règles : pose UNE seule question à la fois, courte et claire. "
        "Ne donne jamais de feedback pendant l'entretien. "
        "Commence par une question d'introduction, puis pose des questions "
        "liées à l'offre (compétences, motivation, mise en situation)."
    )
    st.session_state.messages = [
        {"role": "system", "content": consigne},
        {"role": "user", "content": "Bonjour, je suis prêt pour l'entretien."},
    ]
    question = demander(st.session_state.messages)
    st.session_state.messages.append({"role": "assistant", "content": question})
    st.session_state.feedback = None
    st.session_state.ecran = "entretien"


def transcription():
    """Transforme la conversation en texte simple (sans le message système)."""
    texte = ""
    for m in st.session_state.messages[2:]:
        qui = "Recruteur" if m["role"] == "assistant" else "Candidat"
        texte += f"{qui} : {m['content']}\n\n"
    return texte


def generer_feedback():
    """Demande à Gemma d'évaluer l'entretien (réponse en JSON)."""
    prompt = (
        "Voici un entretien de stage :\n" + transcription() +
        "Évalue le candidat. Réponds UNIQUEMENT en JSON avec ces clés :\n"
        '"score" (nombre entier de 0 à 10),\n'
        '"resume" (résumé de l\'entretien en 3 phrases),\n'
        '"points_forts" (liste de 2 points forts),\n'
        '"clarte" (texte court),\n'
        '"structure" (texte court),\n'
        '"ameliorations" (liste de 3 conseils concrets).'
    )
    brut = demander([{"role": "user", "content": prompt}], en_json=True)
    try:
        st.session_state.feedback = json.loads(brut)
    except json.JSONDecodeError:
        st.session_state.feedback = {"erreur": brut}


def en_liste(valeur):
    """S'assure qu'on a une liste (Gemma renvoie parfois du texte simple)."""
    return valeur if isinstance(valeur, list) else [valeur]


def carte(titre, contenu):
    """Affiche une carte blanche arrondie avec un titre."""
    if isinstance(contenu, list):
        contenu = "<ul>" + "".join(f"<li>{escape(str(c))}</li>" for c in contenu) + "</ul>"
    else:
        contenu = escape(str(contenu))
    st.markdown(f'<div class="carte"><h4>{titre}</h4>{contenu}</div>', unsafe_allow_html=True)


def exporter_texte(fb):
    """Crée le rapport au format texte."""
    lignes = ["RÉSULTATS DE L'ENTRETIEN", "", f"Score : {fb.get('score', '?')}/10", ""]
    lignes += ["RÉSUMÉ", str(fb.get("resume", "")), ""]
    lignes += ["POINTS FORTS"] + [f"- {p}" for p in en_liste(fb.get("points_forts", []))] + [""]
    lignes += ["CLARTÉ", str(fb.get("clarte", "")), ""]
    lignes += ["STRUCTURE", str(fb.get("structure", "")), ""]
    lignes += ["AMÉLIORATIONS"] + [f"- {a}" for a in en_liste(fb.get("ameliorations", []))] + [""]
    lignes += ["TRANSCRIPTION", transcription()]
    return "\n".join(lignes)


def exporter_word(fb):
    """Crée le rapport au format Word (.docx) et renvoie ses octets."""
    doc = Document()
    doc.add_heading("Résultats de l'entretien", 0)
    doc.add_paragraph(f"Score : {fb.get('score', '?')}/10")
    doc.add_heading("Résumé", 1)
    doc.add_paragraph(str(fb.get("resume", "")))
    doc.add_heading("Points forts", 1)
    for p in en_liste(fb.get("points_forts", [])):
        doc.add_paragraph(str(p), style="List Bullet")
    doc.add_heading("Clarté", 1)
    doc.add_paragraph(str(fb.get("clarte", "")))
    doc.add_heading("Structure", 1)
    doc.add_paragraph(str(fb.get("structure", "")))
    doc.add_heading("Améliorations", 1)
    for a in en_liste(fb.get("ameliorations", [])):
        doc.add_paragraph(str(a), style="List Bullet")
    doc.add_heading("Transcription de l'entretien", 1)
    for m in st.session_state.messages[2:]:
        qui = "Recruteur" if m["role"] == "assistant" else "Candidat"
        doc.add_paragraph(f"{qui} : {m['content']}")
    tampon = io.BytesIO()
    doc.save(tampon)
    return tampon.getvalue()

if st.session_state.ecran == "offre":
    st.markdown('<div class="titre">🎤 Coach d\'entretien</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="sous-titre">Colle ou importe une offre de stage, '
        "l'IA joue le recruteur et te donne un feedback.</div>",
        unsafe_allow_html=True,
    )

    fichier_offre = st.file_uploader("📄 Importer l'offre (PDF ou Word)", type=["pdf", "docx"])
    texte_offre = st.text_area("✍️ Ou colle l'offre ici", height=200)

    with st.expander("➕ Ajouter mon CV (optionnel, pour des questions personnalisées)"):
        fichier_cv = st.file_uploader("CV (PDF ou Word)", type=["pdf", "docx"], key="cv")

    if st.button("🚀 Commencer l'entretien", type="primary"):
        offre = lire_fichier(fichier_offre) + "\n" + texte_offre
        if offre.strip():
            with st.spinner("Le recruteur arrive..."):
                demarrer(offre, lire_fichier(fichier_cv))
            st.rerun()
        else:
            st.warning("Ajoute d'abord une offre (fichier ou texte).")

else:
    col_titre, col_quitter = st.columns([4, 1])
    with col_titre:
        st.markdown("## 💬 Entretien en cours")
    with col_quitter:
        st.button("🚪 Quitter", on_click=retour_accueil)

    nb_reponses = sum(1 for m in st.session_state.messages if m["role"] == "user") - 1

    if st.session_state.feedback is None:
        st.progress(min(nb_reponses / NB_QUESTIONS, 1.0), text=f"Question {nb_reponses + 1} sur {NB_QUESTIONS}")

    avatars = {"assistant": "🧑‍💼", "user": "🙂"}
    for m in st.session_state.messages[2:]:
        with st.chat_message(m["role"], avatar=avatars[m["role"]]):
            st.write(m["content"])

    if st.session_state.feedback is None:
        reponse = st.chat_input("Ta réponse...")
        if reponse:
            st.session_state.messages.append({"role": "user", "content": reponse})
            with st.spinner("Le recruteur réfléchit..."):
                if nb_reponses + 1 < NB_QUESTIONS:
                    question = demander(st.session_state.messages)
                    st.session_state.messages.append({"role": "assistant", "content": question})
                else:
                    generer_feedback()
            st.rerun()
    else:
        fb = st.session_state.feedback
        st.markdown("## 📝 Ton feedback")
        if "erreur" in fb:
            st.error("Gemma n'a pas renvoyé un JSON valide. Clique sur Recommencer.")
            st.write(fb["erreur"])
        else:
            try:
                score = int(fb.get("score", 0))
            except (ValueError, TypeError):
                score = 0
            couleur = "#16a34a" if score >= 7 else "#f59e0b" if score >= 5 else "#ef4444"
            st.markdown(
                f'<div class="carte" style="text-align:center">'
                f'<div style="font-size:3.5rem;font-weight:800;color:{couleur}">{score}/10</div>'
                f"<div>Score global</div></div>",
                unsafe_allow_html=True,
            )
            carte("📌 Résumé", fb.get("resume", ""))
            carte("💪 Points forts", en_liste(fb.get("points_forts", [])))
            carte("🔍 Clarté", fb.get("clarte", ""))
            carte("🧩 Structure", fb.get("structure", ""))
            carte("🚀 Améliorations", en_liste(fb.get("ameliorations", [])))

            st.markdown("### 📥 Exporter")
            c1, c2 = st.columns(2)
            with c1:
                st.download_button(
                    "Télécharger en Word", exporter_word(fb), "entretien.docx",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    type="primary",
                )
            with c2:
                st.download_button("Télécharger en texte", exporter_texte(fb), "entretien.txt")

        st.button("🔄 Recommencer", on_click=retour_accueil)
